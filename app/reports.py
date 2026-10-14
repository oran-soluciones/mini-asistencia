"""Reporte de asistencia a Excel (openpyxl).

`resumen_por_dia` es una función pura (se testea sin Excel); `generar_excel` solo escribe.
La regla de la jornada (pares entrada/salida, repetidas, anuladas) vive en app/jornada.py y es
la misma que usan la vista Hoy y la columna Tipo del panel.

Reglas del resumen por día (una fila por empleado y día calendario):
- Las marcaciones válidas van de a pares: 1.ª entrada, 2.ª salida, 3.ª entrada... Cada día
  empieza con una entrada. Horas = suma de los pares completos. No contempla turnos ni horarios.
- Repetidas: una marcación a menos de 2 minutos de la última marcación válida (típico de
  marcar dos veces seguidas) se ignora. Se cuentan en la Observación.
- Anuladas a mano desde el panel: no cuentan (ni en Cantidad); se anotan en la Observación.
- Cantidad impar: al último par le falta la salida y no suma horas. Si ese día es hoy
  (exportación a media jornada) no es un error: dice "Sin salida todavía".
- No se cruza la medianoche: cada día calendario se calcula por separado (un turno nocturno
  aparece como dos días incompletos).

Hojas: "Resumen por día" (una columna Entrada/Salida por par), "Línea de tiempo" (la presencia
de cada día pintada en franjas de 15 minutos), "Detalle" (cada marcación con su Tipo) e "Info".

Cada hoja de datos lleva un encabezado de documento (logo, colegio, título, período y fecha de
emisión) y está configurada para imprimir (A4 horizontal, 1 página de ancho, títulos repetidos).
"""
import logging
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import NamedTuple

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImagenExcel
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.properties import PageSetupProperties
from PIL import Image as ImagenPIL

from app.jornada import REPETIDA_MIN as REPETIDA  # noqa: F401 (nombre histórico en este módulo)
from app.jornada import analizar, anulada, tipos_de_marcaciones

log = logging.getLogger("panel")

# Tope de seguridad de filas del Excel (si hay más se exportan las más nuevas y se avisa en "Info")
LIMITE_EXCEL = 200_000
VERIFICACION_CARA = 15  # deducido, no verificado
FORMATO_FECHA, FORMATO_HORA, FORMATO_HORAS = "dd/mm/yyyy", "hh:mm:ss", "[h]:mm"
NOTAS = [
  "Las marcaciones de cada día van de a pares: 1.ª entrada, 2.ª salida, 3.ª entrada...",
  "Horas = suma de cada par entrada-salida; no contempla turnos ni horarios.",
  "Marcaciones a menos de 2 minutos de la anterior se ignoran.",
  "Las marcaciones anuladas a mano en el panel no cuentan (figuran en Detalle como Anulada).",
  "Cada día se calcula por separado y empieza con una entrada: no se cruza la medianoche.",
]

# Encabezado de documento y estilos (paleta del panel: pizarrón y tiza)
FILA_TITULOS = 5  # filas 1-3: encabezado del documento, 4: espacio, 5: títulos, 6+: datos
PRIMERA_FILA_DATOS = FILA_TITULOS + 1
VERDE, TEXTO, SUAVE = "262626", "1F1F1F", "5C5C5C"
CEBRA, BORDE = "F4F4F4", "DDDDDD"
AMBAR_FONDO, AMBAR_TEXTO = "FBEFD5", "7A4A00"
COLOR_TIPO = {"Entrada": "1F7A4A", "Salida": "4A5A68", "Repetida": "707070", "Anulada": "B42318"}
ALTO_FILA = 20
# Línea de tiempo: franjas de 15 minutos; presencia cerrada, en curso (hoy) y salida faltante
FRANJA_MIN = 15
COLOR_PRESENTE, COLOR_EN_CURSO, COLOR_FALTA = "4A5A68", "1F7A4A", "E3A13B"
LINEA_HORA = "B8B8B8"
LOGO_MAX = (150, 52)  # px: ancho y alto máximos en la hoja (cabe en las columnas A y B)


class Resumen(NamedTuple):
  fecha: datetime  # solo la parte de fecha
  pin: str
  nombre: str
  pares: list[tuple[datetime, datetime | None]]  # (entrada, salida); salida None = falta
  cantidad: int  # marcaciones reales del día (con repetidas, sin anuladas)
  horas: timedelta  # suma de los pares completos
  observacion: str
  alerta: bool = True  # False: la observación es solo informativa (no va en ámbar)

  @property
  def entrada(self) -> datetime:
    return self.pares[0][0]

  @property
  def salida(self) -> datetime | None:
    """Última salida (None si falta la salida del último par)."""
    return self.pares[-1][1]


def _clave_pin(pin: str) -> tuple:
  return (0, int(pin), pin) if pin.isdigit() else (1, 0, pin)


def resumen_por_dia(filas, ahora: datetime | None = None) -> list[Resumen]:
  """`filas`: marcaciones con `fecha_hora` (YYYY-MM-DD HH:MM:SS), `pin`, `nombre` (o None) y
  opcionalmente `anulada`.

  Devuelve una fila por (día, empleado), ordenadas por día y luego PIN (los días con todas sus
  marcaciones anuladas no aparecen). `ahora` (la hora de la exportación) permite distinguir
  "falta la salida" de un día pasado de "sin salida todavía".
  """
  grupos: dict[tuple, tuple[str, list[datetime], list[int]]] = {}
  for f in filas:
    momento = datetime.fromisoformat(f["fecha_hora"])
    _, momentos, anuladas = grupos.setdefault((momento.date(), f["pin"]),
                                              (f["nombre"] or "", [], [0]))
    if anulada(f):
      anuladas[0] += 1
    else:
      momentos.append(momento)
  resultado = []
  for (dia, pin), (nombre, momentos, anuladas) in sorted(
      grupos.items(), key=lambda kv: (kv[0][0], _clave_pin(kv[0][1]))):
    if not momentos:
      continue
    j = analizar(momentos)
    obs, alerta = [], False
    if j.repetidas:
      obs.append(f"{j.repetidas} repetida(s) ignorada(s)")
      alerta = True
    if j.abierta:
      en_curso = ahora is not None and dia == ahora.date()
      obs.append("Sin salida todavía" if en_curso else "Falta la salida (marcaciones impares)")
      alerta = alerta or not en_curso
    if anuladas[0]:
      obs.append(f"{anuladas[0]} anulada(s) a mano")
    resultado.append(Resumen(
      fecha=datetime.combine(dia, datetime.min.time()), pin=pin, nombre=nombre, pares=j.pares,
      cantidad=len(momentos), horas=j.horas, observacion="; ".join(obs), alerta=alerta,
    ))
  return resultado


def _celda_pin(pin: str):
  """Número si es un entero sin ceros a la izquierda; si no, el texto tal cual."""
  return int(pin) if pin.isdigit() and str(int(pin)) == pin else pin


def _dias(horas: timedelta) -> float:
  return horas.total_seconds() / 86400  # duración de Excel: fracción de día


def _logo_excel(ruta: Path | None, tipo: str | None) -> ImagenExcel | None:
  """El logo listo para insertar, o None (sin logo, SVG —Excel no lo admite— o ilegible)."""
  if ruta is None or tipo not in ("image/png", "image/jpeg"):
    return None
  try:
    with ImagenPIL.open(ruta) as original:
      original.load()
      ancho, alto = original.size
      escala = min(LOGO_MAX[0] / ancho, LOGO_MAX[1] / alto, 1)
      mostrado = (max(1, round(ancho * escala)), max(1, round(alto * escala)))
      factor = min(2 * escala, 1)  # al doble de lo que se muestra: nítido al imprimir
      nitido = original.convert("RGBA").resize(
        (max(1, round(ancho * factor)), max(1, round(alto * factor))), ImagenPIL.LANCZOS)
    buffer = BytesIO()
    nitido.save(buffer, format="PNG")
  except (OSError, ValueError, ImagenPIL.DecompressionBombError) as e:
    log.warning("No se pudo usar el logo en el Excel (%s): se genera sin logo", e)
    return None
  imagen = ImagenExcel(buffer)
  imagen.width, imagen.height = mostrado
  imagen.anchor = OneCellAnchor(
    _from=AnchorMarker(col=0, row=0, colOff=pixels_to_EMU(4), rowOff=pixels_to_EMU(8)),
    ext=XDRPositiveSize2D(pixels_to_EMU(mostrado[0]), pixels_to_EMU(mostrado[1])))
  return imagen


def _encabezado_documento(ws, colegio: str, logo: ImagenExcel | None, rango: str) -> None:
  """Filas 1-3: logo, colegio, 'Asistencia del personal' y período / emisión."""
  col = 3 if logo else 1  # el logo ocupa las columnas A y B
  textos = [
    (colegio or "Colegio", Font(size=16, bold=True, color=VERDE), 26),
    ("Asistencia del personal", Font(size=12, bold=True, color=SUAVE), 19),
    (rango, Font(size=10, color=SUAVE), 17),
  ]
  for fila, (texto, fuente, alto) in enumerate(textos, start=1):
    celda = ws.cell(fila, col, texto)
    celda.font = fuente
    celda.alignment = Alignment(vertical="center")
    ws.row_dimensions[fila].height = alto
  ws.row_dimensions[4].height = 8
  if logo:
    ws.add_image(logo)


def _fila_titulos(ws, titulos: list[str], anchos: list[int], centradas: set[int]) -> None:
  """Fila de títulos en verde pizarrón con texto blanco; congela hasta acá."""
  for i, (titulo, ancho) in enumerate(zip(titulos, anchos), start=1):
    celda = ws.cell(FILA_TITULOS, i, titulo)
    celda.font = Font(bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor=VERDE)
    celda.alignment = Alignment(horizontal="center" if i in centradas else "left",
                                vertical="center", indent=0 if i in centradas else 1)
    ws.column_dimensions[get_column_letter(i)].width = ancho
  ws.row_dimensions[FILA_TITULOS].height = 24
  ws.freeze_panes = f"A{PRIMERA_FILA_DATOS}"


def _estilo_fila(ws, fila: int, ncols: int, centradas: set[int]) -> None:
  """Cebra muy suave, línea fina y alineación por columna."""
  relleno = PatternFill("solid", fgColor=CEBRA) if (fila - PRIMERA_FILA_DATOS) % 2 else None
  linea = Border(bottom=Side(style="thin", color=BORDE))
  for i in range(1, ncols + 1):
    celda = ws.cell(fila, i)
    celda.border = linea
    if relleno:
      celda.fill = relleno
    celda.alignment = Alignment(horizontal="center" if i in centradas else "left",
                                vertical="center", indent=0 if i in centradas else 1)
  ws.row_dimensions[fila].height = ALTO_FILA


def _imprimir(ws) -> None:
  """A4 horizontal, 1 página de ancho, títulos repetidos y pie 'Página N de M'."""
  ws.page_setup.orientation = "landscape"
  ws.page_setup.paperSize = ws.PAPERSIZE_A4
  ws.page_setup.fitToWidth = 1
  ws.page_setup.fitToHeight = 0  # las páginas que hagan falta a lo largo
  ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
  ws.print_title_rows = f"{FILA_TITULOS}:{FILA_TITULOS}"
  ws.print_options.horizontalCentered = True
  ws.page_margins.left = ws.page_margins.right = 0.4
  ws.page_margins.top = ws.page_margins.bottom = 0.6
  ws.oddFooter.center.text = "Página &P de &N"
  ws.oddFooter.center.size = 9
  ws.oddFooter.center.color = SUAVE
  ws.sheet_view.showGridLines = False


def _periodo_txt(desde: str | None, hasta: str | None) -> str:
  if desde and hasta:
    return f"Período: {_fecha_txt(desde, '')} al {_fecha_txt(hasta, '')}"
  if desde:
    return f"Desde el {_fecha_txt(desde, '')}"
  if hasta:
    return f"Hasta el {_fecha_txt(hasta, '')}"
  return "Todas las fechas"


def _hm(momento: datetime | None) -> str:
  return f"{momento:%H:%M}" if momento else "?"


def _horario(pares) -> str:
  """'08:00 a 12:00  ·  13:00 a 17:00' ('?' donde falta la salida)."""
  return "  ·  ".join(f"{_hm(e)} a {_hm(s)}" for e, s in pares)


def _rango_horas(resumenes: list[Resumen]) -> tuple[int, int]:
  """De 07 a 19 h, ampliado si alguien marcó antes o después."""
  inicio, fin = 7, 19
  for r in resumenes:
    for e, s in r.pares:
      inicio = min(inicio, e.hour)
      ultimo = s or e
      fin = max(fin, ultimo.hour + (1 if ultimo.minute or ultimo.second else 0))
  return inicio, min(fin, 24)


def _hoja_linea_de_tiempo(ws, resumenes: list[Resumen], generado: datetime, colegio: str,
                          logo: ImagenExcel | None, rango: str) -> None:
  """Una fila por día y empleado; cada columna angosta es una franja de 15 minutos y se pinta
  si la persona estuvo en el colegio en algún momento de esa franja."""
  fijas = ["Fecha", "Nombre", "Horario", "Horas"]
  centradas = {1, 4}
  inicio_h, fin_h = _rango_horas(resumenes)
  por_hora = 60 // FRANJA_MIN
  nfranjas = (fin_h - inicio_h) * por_hora
  primera = len(fijas) + 1  # columna de la primera franja
  paso = timedelta(minutes=FRANJA_MIN)

  _encabezado_documento(ws, colegio, logo, rango)
  _fila_titulos(ws, fijas, [13, 30, 32, 10], centradas)
  for k in range(nfranjas):
    col = primera + k
    ws.column_dimensions[get_column_letter(col)].width = 1.7
    ws.cell(FILA_TITULOS, col).fill = PatternFill("solid", fgColor=VERDE)
  for h in range(fin_h - inicio_h):  # rótulo de cada hora, sobre sus 4 franjas
    col = primera + h * por_hora
    celda = ws.cell(FILA_TITULOS, col, f"{inicio_h + h:02d}")
    celda.font = Font(bold=True, color="FFFFFF", size=9)
    celda.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=FILA_TITULOS, start_column=col, end_row=FILA_TITULOS,
                   end_column=col + por_hora - 1)

  for r in resumenes:
    ws.append([r.fecha.date(), r.nombre or f"PIN {r.pin}", _horario(r.pares), _dias(r.horas)])
    fila = ws.max_row
    _estilo_fila(ws, fila, len(fijas), centradas)
    ws.cell(fila, 1).number_format = FORMATO_FECHA
    ws.cell(fila, 4).number_format = FORMATO_HORAS
    cero = r.fecha + timedelta(hours=inicio_h)
    pintadas: dict[int, str] = {}
    for e, s in r.pares:
      if s:
        fin, color = s, COLOR_PRESENTE
      elif r.fecha.date() == generado.date() and generado > e:
        fin, color = generado, COLOR_EN_CURSO  # sigue en el colegio: hasta la hora de emisión
      else:
        fin, color = e + paso, COLOR_FALTA  # falta la salida: solo se marca la entrada
      for k in range(nfranjas):
        franja = cero + k * paso
        if e < franja + paso and fin > franja:
          pintadas[k] = color
    abajo = Side(style="thin", color=BORDE)
    for k in range(nfranjas):
      celda = ws.cell(fila, primera + k)
      izq = Side(style="thin", color=LINEA_HORA) if k % por_hora == 0 else Side()
      celda.border = Border(left=izq, bottom=abajo)
      if k in pintadas:
        celda.fill = PatternFill("solid", fgColor=pintadas[k])
  ultima_fila = max(ws.max_row, FILA_TITULOS)
  ws.auto_filter.ref = f"A{FILA_TITULOS}:D{ultima_fila}"

  ws.append([])
  for color, texto in [(COLOR_PRESENTE, "En el colegio (entrada y salida marcadas)"),
                       (COLOR_EN_CURSO, "Sigue en el colegio (hoy, hasta la hora de emisión)"),
                       (COLOR_FALTA, "Entrada sin salida (falta marcar la salida)")]:
    ws.append([None, texto])
    ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=color)
    ws.cell(ws.max_row, 2).font = Font(italic=True, color=SUAVE)
  ws.append([None, f"Cada columna angosta son {FRANJA_MIN} minutos; las líneas verticales "
                   "marcan cada hora."])
  ws.cell(ws.max_row, 2).font = Font(italic=True, color=SUAVE)
  _imprimir(ws)


def generar_excel(filas, desde: str | None, hasta: str | None, pin: str | None,
                  nombre_pin: str | None = None, truncado: bool = False,
                  generado: datetime | None = None, colegio: str = "",
                  logo: Path | None = None, tipo_logo: str | None = None) -> bytes:
  """Arma el .xlsx (Resumen por día, Línea de tiempo, Detalle e Info) y lo devuelve en bytes.

  `filas`: marcaciones en orden ascendente con fecha_hora, pin, nombre, verificacion, sn y
  (opcional) anulada. `colegio`/`logo`/`tipo_logo`: nombre y logo cargados en Configuración
  (el SVG no se inserta).
  """
  filas = list(filas)
  generado = generado or datetime.now()
  empleado = "Todos" if not pin else (f"{pin} - {nombre_pin}" if nombre_pin else pin)
  rango = (f"{_periodo_txt(desde, hasta)}  ·  Empleado: {empleado}  ·  "
           f"Emitido el {generado:%d/%m/%Y %H:%M}")
  resumenes = resumen_por_dia(filas, generado)
  wb = Workbook()

  ws = wb.active
  ws.title = "Resumen por día"
  # Columnas: Fecha, PIN, Nombre, (Entrada, Salida) por par, Cantidad, Horas, Observación
  npares = max((len(r.pares) for r in resumenes), default=1)
  if npares == 1:
    titulos_pares = ["Entrada", "Salida"]
  else:
    titulos_pares = [t for n in range(1, npares + 1) for t in (f"Entrada {n}", f"Salida {n}")]
  c_cant = 4 + 2 * npares  # columna de Cantidad; le siguen Horas y Observación
  c_horas, ncols = c_cant + 1, c_cant + 2
  centradas = {1, 2, *range(4, ncols)}
  _encabezado_documento(ws, colegio, _logo_excel(logo, tipo_logo), rango)
  _fila_titulos(ws, ["Fecha", "PIN", "Nombre", *titulos_pares, "Cantidad", "Horas trabajadas",
                     "Observación"],
                [13, 9, 32, *[12] * (2 * npares), 11, 18, 60], centradas)
  total = timedelta(0)
  for r in resumenes:
    total += r.horas
    horas_pares = []
    for e, s in r.pares:
      horas_pares += [e.time(), s.time() if s else None]
    horas_pares += [None] * (2 * npares - len(horas_pares))
    ws.append([r.fecha.date(), _celda_pin(r.pin), r.nombre, *horas_pares, r.cantidad,
               _dias(r.horas), r.observacion])
    fila = ws.max_row
    _estilo_fila(ws, fila, ncols, centradas)
    ws.cell(fila, 1).number_format = FORMATO_FECHA
    for col in range(4, c_cant):
      ws.cell(fila, col).number_format = FORMATO_HORA
    ws.cell(fila, c_horas).number_format = FORMATO_HORAS
    if r.observacion:
      obs = ws.cell(fila, ncols)
      if r.alerta:
        obs.fill = PatternFill("solid", fgColor=AMBAR_FONDO)
        obs.font = Font(color=AMBAR_TEXTO)
      else:
        obs.font = Font(color=SUAVE, italic=True)
      obs.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
  ultima = get_column_letter(ncols)
  ws.auto_filter.ref = f"A{FILA_TITULOS}:{ultima}{max(ws.max_row, FILA_TITULOS)}"  # sin Total
  ws.append(["Total", *[None] * (c_cant - 1), _dias(total)])
  ws.cell(ws.max_row, c_horas).number_format = FORMATO_HORAS
  for celda in ws[ws.max_row]:
    celda.font = Font(bold=True)
    celda.border = Border(top=Side(style="medium", color=VERDE))
  ws.cell(ws.max_row, c_horas).alignment = Alignment(horizontal="center")
  ws.row_dimensions[ws.max_row].height = ALTO_FILA
  ws.append([])
  for nota in NOTAS:
    ws.append([nota])
    ws.cell(ws.max_row, 1).font = Font(italic=True, color=SUAVE)
  _imprimir(ws)

  _hoja_linea_de_tiempo(wb.create_sheet("Línea de tiempo"), resumenes, generado, colegio,
                        _logo_excel(logo, tipo_logo), rango)

  ws = wb.create_sheet("Detalle")
  # Columnas: 1 Fecha, 2 Hora, 3 Tipo, 4 PIN, 5 Nombre, 6 Verificación, 7 SN
  centradas = {1, 2, 3, 4}
  _encabezado_documento(ws, colegio, _logo_excel(logo, tipo_logo), rango)
  _fila_titulos(ws, ["Fecha", "Hora", "Tipo", "PIN", "Nombre", "Verificación", "SN"],
                [13, 11, 14, 9, 32, 15, 22], centradas)
  for f, tipo in zip(filas, tipos_de_marcaciones(filas)):
    momento = datetime.fromisoformat(f["fecha_hora"])
    v = f["verificacion"]
    ws.append([momento.date(), momento.time(), tipo, _celda_pin(f["pin"]), f["nombre"] or "",
               "Cara" if v == VERIFICACION_CARA else ("" if v is None else v), f["sn"]])
    fila = ws.max_row
    _estilo_fila(ws, fila, 7, centradas)
    ws.cell(fila, 1).number_format = FORMATO_FECHA
    ws.cell(fila, 2).number_format = FORMATO_HORA
    ws.cell(fila, 3).font = Font(color=COLOR_TIPO[tipo], italic=tipo in ("Repetida", "Anulada"))
  ws.auto_filter.ref = f"A{FILA_TITULOS}:G{max(ws.max_row, FILA_TITULOS)}"
  _imprimir(ws)

  ws = wb.create_sheet("Info")
  ws.column_dimensions["A"].width = 26
  ws.column_dimensions["B"].width = 90
  info = []
  if colegio:
    info.append(("Colegio", colegio))
  info += [
    ("Desde", _fecha_txt(desde, "sin límite")),
    ("Hasta", _fecha_txt(hasta, "sin límite")),
    ("Empleado", empleado),
    ("Marcaciones exportadas", len(filas)),
    ("Generado", generado.strftime("%d/%m/%Y %H:%M:%S")),
  ]
  if truncado:
    info.append(("Aviso", f"Había más de {LIMITE_EXCEL} marcaciones: se exportaron solo las más "
                          "nuevas. Acotá el rango de fechas o elegí un empleado."))
  info += [("Nota", n) for n in NOTAS]
  for etiqueta, valor in info:
    ws.append([etiqueta, valor])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="left")

  # La hoja que se abre es el resumen por día: primera y activa (solo ella seleccionada)
  wb.active = 0
  for hoja in wb.worksheets:
    hoja.sheet_view.tabSelected = hoja.title == "Resumen por día"

  salida = BytesIO()
  wb.save(salida)
  return salida.getvalue()


def _fecha_txt(iso: str | None, vacio: str) -> str:
  return datetime.fromisoformat(iso).strftime("%d/%m/%Y") if iso else vacio
