from datetime import date, datetime, time, timedelta
from io import BytesIO

from openpyxl import load_workbook
from PIL import Image

from app import colegio, db, panel, reports
from app.iclock import parse_users
from conftest import fixture_text
from test_panel import SN, cargar, con_datos, marcacion


def fila(pin, fecha_hora, nombre="Ana", verificacion=15):
  return {"pin": pin, "fecha_hora": fecha_hora, "nombre": nombre,
          "verificacion": verificacion, "sn": SN}


def horas(r):
  return r.horas.total_seconds() / 3600


def test_resumen_par_simple():
  [r] = reports.resumen_por_dia([fila("7", "2026-10-12 08:00:00"), fila("7", "2026-10-12 17:00:00")])
  assert (r.pin, r.nombre, r.cantidad) == ("7", "Ana", 2)
  assert r.fecha.date() == date(2026, 10, 12)
  assert r.entrada == datetime(2026, 10, 12, 8) and r.salida == datetime(2026, 10, 12, 17)
  assert horas(r) == 9 and r.observacion == ""


def test_resumen_una_sola_marcacion():
  [r] = reports.resumen_por_dia([fila("7", "2026-10-12 08:00:00")])
  assert r.entrada == datetime(2026, 10, 12, 8) and r.salida is None
  assert r.horas == timedelta(0) and r.cantidad == 1
  assert r.observacion == "Falta la salida (marcaciones impares)"


def test_resumen_repetidas_a_menos_de_2_minutos():
  [r] = reports.resumen_por_dia([
    fila("7", "2026-10-12 08:00:00"), fila("7", "2026-10-12 08:01:59"),
    fila("7", "2026-10-12 17:00:00"), fila("7", "2026-10-12 17:00:30"),
  ])
  assert r.cantidad == 4 and horas(r) == 9
  assert r.entrada == datetime(2026, 10, 12, 8) and r.salida == datetime(2026, 10, 12, 17)
  assert r.observacion == "2 repetida(s) ignorada(s)"
  # justo 2 minutos ya no es repetida
  [r] = reports.resumen_por_dia([fila("7", "2026-10-12 08:00:00"), fila("7", "2026-10-12 08:02:00")])
  assert r.observacion == "" and r.horas == timedelta(minutes=2)


def test_resumen_solo_repetidas_queda_sin_salida():
  [r] = reports.resumen_por_dia([fila("7", "2026-10-12 08:00:00"), fila("7", "2026-10-12 08:00:40")])
  assert r.salida is None and r.horas == timedelta(0) and r.cantidad == 2
  assert r.observacion == "1 repetida(s) ignorada(s); Falta la salida (marcaciones impares)"


def test_resumen_de_a_pares_suma_cada_par():
  [r] = reports.resumen_por_dia([
    fila("7", "2026-10-12 17:00:00"), fila("7", "2026-10-12 08:00:00"),  # desordenadas
    fila("7", "2026-10-12 12:00:00"), fila("7", "2026-10-12 13:00:00"),
  ])
  assert r.cantidad == 4 and horas(r) == 8 and r.observacion == ""  # 4 h + 4 h
  assert [(e.hour, s.hour) for e, s in r.pares] == [(8, 12), (13, 17)]


def test_resumen_impar_el_ultimo_par_no_suma():
  [r] = reports.resumen_por_dia([fila("7", "2026-10-12 08:00:00"), fila("7", "2026-10-12 12:00:00"),
                                 fila("7", "2026-10-12 13:00:00")])
  assert horas(r) == 4 and r.salida is None and r.alerta
  assert r.observacion == "Falta la salida (marcaciones impares)"


def test_resumen_anuladas_no_cuentan_y_se_anotan():
  filas = [fila("7", "2026-10-12 08:00:00"), fila("7", "2026-10-12 08:05:00"),
           fila("7", "2026-10-12 17:00:00")]
  filas[1]["anulada"] = "2026-10-12 18:00:00"
  [r] = reports.resumen_por_dia(filas)
  assert horas(r) == 9 and r.cantidad == 2 and not r.alerta
  assert r.observacion == "1 anulada(s) a mano"
  # un día con todas sus marcaciones anuladas no aparece
  filas = [dict(fila("7", "2026-10-12 08:00:00"), anulada="x")]
  assert reports.resumen_por_dia(filas) == []


def test_resumen_dos_empleados_y_dos_dias_ordenados_por_dia_y_pin():
  res = reports.resumen_por_dia([
    fila("10", "2026-10-11 08:00:00", "Zoe"), fila("10", "2026-10-11 16:00:00", "Zoe"),
    fila("9", "2026-10-12 09:00:00", "Ana"), fila("9", "2026-10-12 12:30:00", "Ana"),
    fila("9", "2026-10-11 08:00:00", "Ana"), fila("9", "2026-10-11 17:00:00", "Ana"),
    fila("10", "2026-10-12 10:00:00", "Zoe"),
  ])
  assert [(r.fecha.date().day, r.pin) for r in res] == [(11, "9"), (11, "10"), (12, "9"), (12, "10")]
  assert [horas(r) for r in res] == [9, 8, 3.5, 0]


def test_resumen_pin_sin_nombre():
  [r] = reports.resumen_por_dia([fila("99", "2026-10-12 08:00:00", None)])
  assert r.nombre == ""


def test_resumen_no_cruza_la_medianoche():
  res = reports.resumen_por_dia([fila("7", "2026-10-11 22:00:00"), fila("7", "2026-10-12 06:00:00")])
  assert len(res) == 2 and all(r.salida is None for r in res)


# --- Endpoint -----------------------------------------------------------------------

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def abrir(r):
  return load_workbook(BytesIO(r.content))


def datos(ws):
  """Filas de datos (debajo del encabezado del documento y de la fila de títulos)."""
  return [[c.value for c in row] for row in ws.iter_rows(min_row=reports.PRIMERA_FILA_DATOS)]


def titulos(ws):
  return [c.value for c in ws[reports.FILA_TITULOS]]


def test_excel_respuesta_y_hojas(client, db_path):
  con_datos(db_path)
  r = client.get("/marcaciones/excel?desde=&hasta=")
  assert r.status_code == 200
  assert r.headers["content-type"] == XLSX
  assert r.headers["content-disposition"] == 'attachment; filename="asistencia_todo_todo.xlsx"'
  wb = abrir(r)
  assert wb.sheetnames == ["Resumen por día", "Línea de tiempo", "Detalle", "Info"]  # se abre en el resumen
  assert wb.active.title == "Resumen por día"
  assert [bool(h.sheet_view.tabSelected) for h in wb.worksheets] == [True, False, False, False]
  ws = wb["Detalle"]
  assert titulos(ws) == ["Fecha", "Hora", "Tipo", "PIN", "Nombre", "Verificación", "SN"]
  # Los títulos están en la fila 5 (debajo del encabezado del documento): se congelan y filtran
  assert ws.freeze_panes == "A6" and ws.auto_filter.ref == "A5:G10"
  assert ws["A5"].font.b and ws["A5"].fill.fgColor.rgb.endswith("262626")
  assert ws["A5"].font.color.rgb.endswith("FFFFFF")
  filas = datos(ws)
  assert len(filas) == 5
  # Orden ascendente; Fecha y Hora son valores reales con formato
  assert filas[0][0] == datetime(2026, 10, 10) and filas[0][1] == time(8, 0, 0)
  assert filas[-1][3:] == [99, None, "Cara", SN]  # PIN sin empleado: nombre vacío
  assert ws["A6"].number_format == "dd/mm/yyyy" and ws["B6"].number_format == "hh:mm:ss"
  assert [f[0] for f in filas] == sorted(f[0] for f in filas)
  assert filas[0][4] == "Ana Gómez" and filas[0][5] == "Cara"
  assert [f[5] for f in filas].count(1) == 1  # verificación distinta de 15: el número


def test_excel_resumen_horas_y_total(client, db_path, monkeypatch):
  monkeypatch.setattr(db, "ahora", lambda: datetime(2026, 10, 20, 9, 0))  # los días ya pasaron
  db.upsert_empleados(SN, parse_users(fixture_text("operlog_usuarios.txt")), db_path)
  cargar(db_path, [
    marcacion(7, "2026-10-12 08:00:00"), marcacion(7, "2026-10-12 17:00:00"),
    marcacion(8, "2026-10-12 09:00:00"), marcacion(8, "2026-10-12 09:01:00"),
    marcacion(8, "2026-10-12 12:00:00"), marcacion(8, "2026-10-12 13:30:00"),
    marcacion(8, "2026-10-12 15:00:00"),
    marcacion(99, "2026-10-13 10:00:00"),
  ])
  ws = abrir(client.get("/marcaciones/excel?desde=&hasta=&pin="))["Resumen por día"]
  # Una columna Entrada/Salida por par (Luis salió y volvió: 2 pares)
  assert titulos(ws) == [
    "Fecha", "PIN", "Nombre", "Entrada 1", "Salida 1", "Entrada 2", "Salida 2", "Cantidad",
    "Horas trabajadas", "Observación"]
  assert ws.freeze_panes == "A6" and ws.auto_filter.ref == "A5:J8"
  [a, b, c] = datos(ws)[:3]
  assert a[:8] == [datetime(2026, 10, 12), 7, "Ana Gómez", time(8), time(17), None, None, 2]
  assert a[8] == timedelta(hours=9) and a[9] is None
  assert b[1:8] == [8, "Luis Pérez", time(9), time(12), time(13, 30), time(15), 5]
  assert b[8] == timedelta(hours=4.5) and b[9] == "1 repetida(s) ignorada(s)"
  assert c[1:8] == [99, None, time(10), None, None, None, 1]
  assert c[8] == timedelta(0) and "Falta la salida" in c[9]
  assert ws["I6"].number_format == "[h]:mm" and ws.column_dimensions["C"].width >= 30
  assert ws["J8"].fill.fgColor.rgb.endswith("FBEFD5") and not ws["J6"].fill.fgColor.rgb.endswith("FBEFD5")
  assert ws["A7"].fill.fgColor.rgb.endswith("F4F4F4") and ws["A6"].fill.fill_type is None  # cebra
  total = [row for row in ws.iter_rows() if row[0].value == "Total"][0]
  assert total[8].value == timedelta(hours=13.5) and total[8].number_format == "[h]:mm"
  notas = " ".join(str(row[0].value) for row in ws.iter_rows() if row[0].value)
  assert "van de a pares" in notas and "no contempla turnos ni horarios" in notas


def test_excel_linea_de_tiempo(client, db_path, monkeypatch):
  monkeypatch.setattr(db, "ahora", lambda: datetime(2026, 10, 20, 9, 0))
  cargar(db_path, [
    marcacion(7, "2026-10-12 08:00:00"), marcacion(7, "2026-10-12 12:00:00"),
    marcacion(7, "2026-10-12 13:00:00"), marcacion(7, "2026-10-12 17:00:00"),
    marcacion(8, "2026-10-12 09:00:00"),
  ])
  ws = abrir(client.get("/marcaciones/excel?desde=&hasta="))["Línea de tiempo"]
  t = titulos(ws)
  assert t[:4] == ["Fecha", "Nombre", "Horario", "Horas"] and t[4] == "07"  # de 07 a 19 h
  assert len(t) == 4 + 12 * 4  # franjas de 15 minutos
  [a, b] = datos(ws)[:2]
  assert a[2] == "08:00 a 12:00  ·  13:00 a 17:00" and a[3] == timedelta(hours=8)
  assert b[2] == "09:00 a ?"

  def color(fila, hora, minuto=0):
    col = 5 + (hora - 7) * 4 + minuto // 15
    celda = ws.cell(fila, col)
    return celda.fill.fgColor.rgb[-6:] if celda.fill.fill_type else None

  assert color(6, 7, 45) is None and color(6, 8) == reports.COLOR_PRESENTE
  assert color(6, 11, 45) == reports.COLOR_PRESENTE and color(6, 12) is None  # el hueco
  assert color(6, 13) == reports.COLOR_PRESENTE and color(6, 17) is None
  assert color(7, 9) == reports.COLOR_FALTA and color(7, 9, 15) is None  # sin salida


def test_excel_info(client, db_path):
  con_datos(db_path)
  ws = abrir(client.get("/marcaciones/excel?desde=2026-10-11&hasta=2026-10-12&pin=8"))["Info"]
  info = {row[0].value: row[1].value for row in ws.iter_rows()}
  assert info["Desde"] == "11/10/2026" and info["Hasta"] == "12/10/2026"
  assert info["Empleado"] == "8 - Luis Pérez" and info["Marcaciones exportadas"] == 2
  assert info["Generado"]


def test_excel_respeta_filtros(client, db_path):
  con_datos(db_path)
  r = client.get("/marcaciones/excel?desde=2026-10-11&hasta=2026-10-11")
  assert r.headers["content-disposition"] == \
    'attachment; filename="asistencia_2026-10-11_2026-10-11.xlsx"'
  assert len(datos(abrir(r)["Detalle"])) == 2
  assert len(datos(abrir(client.get("/marcaciones/excel?desde=2026-10-12&hasta="))["Detalle"])) == 2
  r = client.get("/marcaciones/excel?desde=&hasta=2026-10-10")
  assert len(datos(abrir(r)["Detalle"])) == 1
  assert 'filename="asistencia_todo_2026-10-10.xlsx"' in r.headers["content-disposition"]
  filas = datos(abrir(client.get("/marcaciones/excel?desde=&hasta=&pin=8"))["Detalle"])
  assert [f[3] for f in filas] == [8, 8]


def test_excel_sin_filtros_es_hoy(client, db_path):
  hoy = date.today().isoformat()
  cargar(db_path, [marcacion(7, f"{hoy} 07:30:00"), marcacion(7, "2026-01-01 07:30:00")])
  r = client.get("/marcaciones/excel")
  assert f"asistencia_{hoy}_{hoy}.xlsx" in r.headers["content-disposition"]
  assert len(datos(abrir(r)["Detalle"])) == 1


def test_excel_sin_limite_de_la_tabla_y_con_tope(client, db_path, monkeypatch):
  monkeypatch.setattr(panel, "LIMITE_FILAS", 2)
  con_datos(db_path)
  wb = abrir(client.get("/marcaciones/excel?desde=&hasta="))
  assert len(datos(wb["Detalle"])) == 5
  monkeypatch.setattr(reports, "LIMITE_EXCEL", 3)
  wb = abrir(client.get("/marcaciones/excel?desde=&hasta="))
  filas = datos(wb["Detalle"])
  assert len(filas) == 3 and filas[0][0] == datetime(2026, 10, 11)  # las 3 más nuevas
  assert any(row[0].value == "Aviso" for row in wb["Info"].iter_rows())


def test_excel_vacio(client, db_path):
  wb = abrir(client.get("/marcaciones/excel?desde=&hasta="))
  assert datos(wb["Detalle"]) == []
  assert wb["Resumen por día"]["A6"].value == "Total"


def test_boton_exportar_en_la_pagina(client, db_path):
  html = client.get("/marcaciones").text
  assert 'formaction="/marcaciones/excel"' in html and "Exportar a Excel" in html


def test_excel_rechaza_clientes_no_locales(client_externo, db_path):
  assert client_externo.get("/marcaciones/excel").status_code == 403
  assert not db_path.exists()


# --- Favicon ------------------------------------------------------------------------

def test_favicon_no_crea_sample(client, tmp_path):
  r = client.get("/favicon.ico")
  assert r.status_code == 204
  assert not list(tmp_path.glob("*.txt"))
  assert 'href="/static/icono.svg"' in client.get("/marcaciones").text
  assert client.get("/static/icono.svg").status_code == 200


# --- Encabezado de documento, logo, Tipo e impresión ---------------------------------------

EMISION = datetime(2026, 10, 13, 14, 5)
COLEGIO = "E.E.T. N° 8.102 San José Obrero"


def png(ruta, tamano=(500, 500)):
  imagen = Image.new("RGBA", tamano, (0, 0, 0, 0))
  imagen.paste((0, 0, 0, 255), (10, 10, tamano[0] - 10, tamano[1] - 10))
  imagen.save(ruta)
  return ruta


def excel(filas=None, **kwargs):
  filas = filas or [fila("7", "2026-10-12 08:00:00"), fila("7", "2026-10-12 17:00:00")]
  kwargs.setdefault("generado", EMISION)
  contenido = reports.generar_excel(filas, "2026-10-11", "2026-10-12", None, **kwargs)
  return load_workbook(BytesIO(contenido))


def test_excel_encabezado_sin_logo_ni_nombre():
  wb = excel()
  for hoja in ("Detalle", "Resumen por día"):
    ws = wb[hoja]
    assert ws["A1"].value == "Colegio" and ws["A2"].value == "Asistencia del personal"
    assert "Período: 11/10/2026 al 12/10/2026" in ws["A3"].value
    assert "Empleado: Todos" in ws["A3"].value and "Emitido el 13/10/2026 14:05" in ws["A3"].value
    assert ws["A1"].font.b and ws._images == []
  assert "Colegio" not in [row[0].value for row in wb["Info"].iter_rows()]


def test_excel_encabezado_con_nombre_y_logo(tmp_path):
  logo = png(tmp_path / "logo.png")
  wb = excel(colegio=COLEGIO, logo=logo, tipo_logo="image/png")
  for hoja in ("Detalle", "Resumen por día"):
    ws = wb[hoja]
    assert ws["C1"].value == COLEGIO and ws["A1"].value is None  # A y B son del logo
    assert ws["C2"].value == "Asistencia del personal"
    assert len(ws._images) == 1
    ancho, alto = ws._images[0].anchor.ext.width, ws._images[0].anchor.ext.height
    assert ancho <= 150 * 9525 and alto <= 52 * 9525  # entra en las columnas A y B
    assert abs(ancho - alto) <= 9525  # el emblema es cuadrado: conserva la proporción
  info = {row[0].value: row[1].value for row in wb["Info"].iter_rows()}
  assert info["Colegio"] == COLEGIO


def test_excel_logo_chico_no_se_agranda(tmp_path):
  logo = png(tmp_path / "logo.png", (40, 30))
  ws = excel(logo=logo, tipo_logo="image/png")["Detalle"]
  assert ws._images[0].anchor.ext.width == 40 * 9525


def test_excel_svg_o_logo_ilegible_se_omiten_sin_romper(tmp_path):
  svg = tmp_path / "logo.svg"
  svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1 1"/>', encoding="utf-8")
  ws = excel(colegio=COLEGIO, logo=svg, tipo_logo="image/svg+xml")["Detalle"]
  assert ws._images == [] and ws["A1"].value == COLEGIO
  roto = tmp_path / "roto.png"
  roto.write_bytes(b"no es una imagen")
  ws = excel(colegio=COLEGIO, logo=roto, tipo_logo="image/png")["Detalle"]
  assert ws._images == [] and ws["A1"].value == COLEGIO
  ws = excel(colegio=COLEGIO, logo=tmp_path / "no_existe.png", tipo_logo="image/png")["Detalle"]
  assert ws._images == []


def test_excel_endpoint_usa_nombre_y_logo_de_configuracion(client, db_path, tmp_path):
  con_datos(db_path)
  colegio.guardar(COLEGIO, ("png", png(tmp_path / "x.png").read_bytes()))
  ws = abrir(client.get("/marcaciones/excel?desde=&hasta="))["Detalle"]
  assert ws["C1"].value == COLEGIO and len(ws._images) == 1
  colegio.quitar_logo()
  colegio.guardar("")
  ws = abrir(client.get("/marcaciones/excel?desde=&hasta="))["Detalle"]
  assert ws["A1"].value == "Colegio" and ws._images == []


def test_excel_columna_tipo_en_detalle():
  wb = excel([
    fila("7", "2026-10-12 08:00:00"), fila("7", "2026-10-12 08:01:00"),
    fila("7", "2026-10-12 12:00:00"), fila("7", "2026-10-12 17:00:00"),
    fila("8", "2026-10-12 09:00:00", "Luis"),
  ])
  ws = wb["Detalle"]
  assert [f[2] for f in datos(ws)] == ["Entrada", "Repetida", "Salida", "Entrada", "Entrada"]
  assert ws["C7"].font.i and not ws["C6"].font.i  # la repetida va en cursiva y apagada


def test_excel_listo_para_imprimir():
  wb = excel()
  for hoja in ("Detalle", "Resumen por día"):
    ws = wb[hoja]
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.paperSize == 9  # A4
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
    assert ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.print_title_rows == "$5:$5"
    assert ws.oddFooter.center.text == "Página &P de &N"


def test_excel_total_en_negrita_y_tipos_reales():
  ws = excel()["Resumen por día"]
  total = [row for row in ws.iter_rows() if row[0].value == "Total"][0]
  assert total[0].font.b and total[6].font.b
  assert total[6].value == timedelta(hours=9) and total[6].number_format == "[h]:mm"
  assert ws["A6"].value == datetime(2026, 10, 12) and ws["D6"].value == time(8)


# --- Hoy: "Sin salida todavía" vs "Falta la salida" -------------------------------------------

def test_resumen_del_dia_de_hoy_con_una_marcacion_no_es_un_error():
  hoy = datetime(2026, 10, 13, 11, 0)
  res = reports.resumen_por_dia([
    fila("7", "2026-10-13 08:00:00"),  # hoy: sigue en el colegio
    fila("8", "2026-10-12 08:00:00"),  # ayer: se olvidó de marcar la salida
    fila("9", "2026-10-13 08:00:00"), fila("9", "2026-10-13 08:00:30"),  # hoy, con repetida
  ], hoy)
  por_pin = {r.pin: r for r in res}
  assert por_pin["7"].observacion == "Sin salida todavía" and not por_pin["7"].alerta
  assert por_pin["8"].observacion == "Falta la salida (marcaciones impares)" and por_pin["8"].alerta
  assert por_pin["9"].observacion == "1 repetida(s) ignorada(s); Sin salida todavía"
  assert por_pin["9"].alerta  # la repetida sí es una observación a mirar
  # sin "ahora" (los usos de siempre) todo día con una sola marcación es un error
  assert reports.resumen_por_dia([fila("7", "2026-10-13 08:00:00")])[0].alerta


def test_excel_a_media_jornada_no_pinta_en_ambar_lo_de_hoy():
  ws = excel([fila("7", "2026-10-13 08:00:00"), fila("8", "2026-10-12 08:00:00", "Luis")],
             generado=datetime(2026, 10, 13, 11, 0))["Resumen por día"]
  ayer, hoy = datos(ws)[0], datos(ws)[1]
  assert ayer[7] == "Falta la salida (marcaciones impares)" and hoy[7] == "Sin salida todavía"
  assert ws["H6"].fill.fgColor.rgb.endswith("FBEFD5")      # el día pasado, en ámbar
  assert not ws["H7"].fill.fgColor.rgb.endswith("FBEFD5")  # hoy: informativo, sin ámbar
  assert ws["H7"].font.i


def test_excel_endpoint_usa_la_hora_de_la_pc_para_decidir_que_es_hoy(client, db_path, monkeypatch):
  monkeypatch.setattr(db, "ahora", lambda: datetime(2026, 10, 12, 10, 0))
  cargar(db_path, [marcacion(7, "2026-10-12 08:00:00")])
  ws = abrir(client.get("/marcaciones/excel?desde=&hasta="))["Resumen por día"]
  assert datos(ws)[0][7] == "Sin salida todavía"
  monkeypatch.setattr(db, "ahora", lambda: datetime(2026, 10, 13, 10, 0))
  ws = abrir(client.get("/marcaciones/excel?desde=&hasta="))["Resumen por día"]
  assert datos(ws)[0][7] == "Falta la salida (marcaciones impares)"
